import io
from collections.abc import Callable
from datetime import date
from types import SimpleNamespace
from typing import Any

import openpyxl
import pytest
from accounts.models import User
from django.contrib.auth.models import Group, Permission
from django.test import Client
from django.utils import timezone
from employees.models import Employee

from reports.excel import workbook_bytes
from reports.filters import Filters
from reports.registry import REPORTS
from reports.table import Cell, Table
from reports.types import Column, Report

pytestmark = pytest.mark.django_db

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _sheet(data: bytes) -> Any:
    """The first sheet of an .xlsx file. Untyped on purpose: openpyxl's stubs type `active` as Optional."""
    return openpyxl.load_workbook(io.BytesIO(data)).active


def _remove_export_switch() -> None:
    """Take `accounts.export_lists` away from the HR group (the accounts migrations give it to HR)."""
    export = Permission.objects.get(codename="export_lists", content_type__app_label="accounts")
    Group.objects.get(name="HR").permissions.remove(export)


def test_excel_matches_screen(client: Client, login_hr: User, hired: Employee) -> None:
    resp = client.get("/reports/staff-list/excel/?cols=code&cols=first_hire_date")
    ws = _sheet(resp.content)
    assert ws["A1"].value == "Ажилчдын жагсаалт"
    assert [c.value for c in ws[4]] == ["Код", "Ажилд орсон"]
    assert ws["B5"].value.date() == date(2026, 1, 5) and ws["B5"].number_format == "yyyy-mm-dd"


def test_excel_needs_export_switch(client: Client, login_as: Callable[..., User]) -> None:
    _remove_export_switch()
    login_as(client, "HR")
    assert client.get("/reports/staff-list/excel/").status_code == 403


def test_excel_drops_forbidden_columns(client: Client, login_hr: User) -> None:
    resp = client.get("/reports/staff-list/excel/?cols=code&cols=salary")
    ws = _sheet(resp.content)
    assert [c.value for c in ws[4] if c.value is not None] == ["Код"]


def test_excel_response_is_an_attachment_named_by_key_and_day(client: Client, login_hr: User) -> None:
    resp = client.get("/reports/staff-list/excel/?cols=code")
    assert resp.status_code == 200 and resp["Content-Type"] == XLSX
    assert resp["Content-Disposition"] == (f'attachment; filename="staff-list-{timezone.localdate().isoformat()}.xlsx"')


def test_excel_unknown_report_is_404(client: Client, login_hr: User) -> None:
    assert client.get("/reports/no-such-report/excel/").status_code == 404


def test_excel_button_only_with_export_switch(client: Client, login_as: Callable[..., User]) -> None:
    login_as(client, "HR", "hr-with-export")
    assert 'href="/reports/staff-list/excel/' in client.get("/reports/staff-list/").content.decode()
    _remove_export_switch()
    login_as(client, "HR", "hr-without-export")
    assert "/excel/" not in client.get("/reports/staff-list/").content.decode()


def _cell(value: Any) -> Cell:
    return Cell(value=value)


def test_workbook_layout_groups_and_total() -> None:
    columns = [
        Column(key="name", label="Нэр", value=lambda r: r),
        Column(key="pay", label="Цалин", value=lambda r: r, numeric=True),
        Column(key="day", label="Огноо", value=lambda r: r),
    ]
    table = Table(
        columns=columns,
        groups=[
            (
                "Санхүү",
                [[_cell("Бат"), _cell(1000), _cell(date(2026, 1, 5))], [_cell("Дорж"), _cell(2000), _cell(None)]],
            ),
            ("Хэлтэсгүй", [[_cell("Сараа"), _cell(None), _cell(None)]]),
        ],
        count=3,
        filters=Filters(group="department"),
    )
    data = workbook_bytes("Туршилт", ["Төлөв: Бүгд", "Хэлтэс: Санхүү"], table, sheet="fake")
    ws = _sheet(data)

    assert ws.title == "fake"
    assert ws["A1"].value == "Туршилт" and ws["A1"].font.bold
    assert ws["A2"].value == "Төлөв: Бүгд; Хэлтэс: Санхүү"
    assert [c.value for c in ws[4]] == ["Нэр", "Цалин", "Огноо"] and ws["A4"].font.bold
    assert ws["A5"].value == "Санхүү · 2" and ws["A5"].font.bold
    assert [c.value for c in ws[6]][:2] == ["Бат", 1000] and ws["C6"].value.date() == date(2026, 1, 5)
    assert ws["C6"].number_format == "yyyy-mm-dd"
    assert [c.value for c in ws[7]] == ["Дорж", 2000, None]
    assert ws["A8"].value == "Хэлтэсгүй · 1" and ws["A8"].font.bold
    assert [c.value for c in ws[9]] == ["Сараа", None, None]
    assert ws["A10"].value == "Нийт: 3"


def test_workbook_without_group_has_no_group_rows_and_caps_width() -> None:
    long_text = "а" * 80
    table = Table(
        columns=[Column(key="name", label="Нэр", value=lambda r: r)],
        groups=[("", [[_cell(long_text)], [_cell("Бат")]])],
        count=2,
        filters=Filters(),
    )
    ws = _sheet(workbook_bytes("Туршилт", [], table, sheet="fake"))
    assert [c.value for c in ws[5]] == [long_text] and ws["A6"].value == "Бат"
    assert ws["A7"].value == "Нийт: 2"
    assert ws.column_dimensions["A"].width == 50


def test_text_starting_with_equals_is_not_a_formula() -> None:
    table = Table(
        columns=[Column(key="name", label="Нэр", value=lambda r: r)],
        groups=[("", [[_cell('=HYPERLINK("http://x")')]])],
        count=1,
        filters=Filters(),
    )
    ws = _sheet(workbook_bytes("Туршилт", [], table, sheet="fake"))
    assert ws["A5"].data_type == "s" and ws["A5"].value == '=HYPERLINK("http://x")'


def test_workbook_takes_three_arguments_with_a_default_sheet_name() -> None:
    table = Table(
        columns=[Column(key="name", label="Нэр", value=lambda r: r)], groups=[("", [])], count=0, filters=Filters()
    )
    ws = _sheet(workbook_bytes("Туршилт", [], table))
    assert ws.title == "Тайлан" and ws["A1"].value == "Туршилт"


def test_excel_drops_personal_columns_for_accountant(client: Client, login_as: Callable[..., User]) -> None:
    login_as(client, "Accountant")
    resp = client.get("/reports/staff-list/excel/?cols=code&cols=clan_name&cols=gender&cols=birthplace")
    assert resp.status_code == 200
    assert [c.value for c in _sheet(resp.content)[4] if c.value is not None] == ["Код"]


def test_control_characters_are_dropped_from_text_cells(
    client: Client, login_hr: User, monkeypatch: pytest.MonkeyPatch
) -> None:
    report = Report(
        key="ctl",
        title="Тэмдэгт",
        filters=(),
        columns=(Column(key="v", label="Утга", value=lambda r: r.v),),
        default_columns=("v",),
        rows=lambda user, filters: [SimpleNamespace(v="AB")],  # a control character a MACS text may hold
    )
    monkeypatch.setitem(REPORTS, "ctl", report)
    resp = client.get("/reports/ctl/excel/")
    assert resp.status_code == 200
    assert _sheet(resp.content)["A5"].value == "AB"


def test_workbook_without_count_has_no_total_row() -> None:
    table = Table(
        columns=[Column(key="name", label="Нэр", value=lambda r: r)],
        groups=[("", [[_cell("Бат")]])],
        count=1,
        filters=Filters(),
        shows_count=False,
    )
    ws = _sheet(workbook_bytes("Туршилт", [], table))
    assert ws["A5"].value == "Бат" and ws["A6"].value is None
