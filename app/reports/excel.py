"""The Excel file of a report: the table the screen and the print show, with title, filter lines and total."""

import io
from datetime import date
from decimal import Decimal
from typing import cast

from openpyxl import Workbook
from openpyxl.cell import Cell as XlCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from reports.table import Table

MAX_WIDTH = 50  # characters; a wider column is capped so one long text does not make the sheet unreadable
DATE_FORMAT = "yyyy-mm-dd"
BOLD = Font(bold=True)


def _text(value: object) -> str:
    """The text a value takes in the sheet, used to size the columns."""
    if value is None:
        return ""
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def _write(ws: Worksheet, row: int, column: int, value: object, bold: bool = False) -> None:
    """Writes one value keeping its type: numbers stay numbers, dates stay dates, blanks stay empty.
    Text is always stored as text, so a value starting with "=" never becomes a formula."""
    cell = cast(XlCell, ws.cell(row=row, column=column))  # a sheet written here has no merged cells
    if bold:
        cell.font = BOLD
    if value is None:
        return
    if isinstance(value, date):
        cell.value = value
        cell.number_format = DATE_FORMAT
    elif isinstance(value, (int, float, Decimal)):
        cell.value = value
    else:
        cell.value = ILLEGAL_CHARACTERS_RE.sub("", _text(value))  # control characters old MACS text may hold
        cell.data_type = "s"


def workbook_bytes(title: str, lines: list[str], table: Table, *, sheet: str = "Тайлан") -> bytes:
    """The report as an .xlsx file. Row 1 the title, row 2 the filter lines, row 4 the headers, then the rows.
    When the table is grouped, a bold «label · N» row starts each group. The last row is the total.
    `sheet` names the sheet; the views pass the report key."""
    wb = Workbook()
    ws = cast(Worksheet, wb.active)  # a new workbook always has its first sheet
    ws.title = sheet[:31]  # Excel's sheet-name limit
    _write(ws, 1, 1, title, bold=True)
    _write(ws, 2, 1, "; ".join(lines) or None)

    headers = [c.label for c in table.columns]
    widths = [len(h) for h in headers]
    for col, label in enumerate(headers, start=1):
        _write(ws, 4, col, label, bold=True)

    row = 5
    for label, rows in table.groups:
        if table.filters.group:
            _write(ws, row, 1, f"{label} · {len(rows)}", bold=True)
            row += 1
        for values in rows:
            for col, cell in enumerate(values, start=1):
                _write(ws, row, col, cell.value)
                widths[col - 1] = max(widths[col - 1], len(_text(cell.value)))
            row += 1
    if table.shows_count:
        _write(ws, row, 1, f"Нийт: {table.count}")

    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = min(width + 2, MAX_WIDTH)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
